"""Stage 7: Validation Node (TRD Section 11.3, Table 30, Table 44, Section 21, Section 22)."""

import logging
from pathlib import Path
from typing import Any, Dict, List, Optional
import docx
import fitz  # PyMuPDF
import openpyxl
import pptx

from backend.app.agent.event_broadcaster import get_event_broadcaster
from backend.app.agent.nodes.tool_selection import is_approval_note_goal
from backend.app.agent.state import AgentState
from backend.app.persistence.artifact_repository import ArtifactRepository
from backend.app.persistence.db import get_db_context

logger = logging.getLogger("sovereign_workbench.agent.node.validation")


def validation_node(state: AgentState) -> Dict[str, Any]:
    """Validation node: verifies plan execution, tool outputs, artifacts, and tests."""
    task_id = state["task_id"]
    task_type = state.get("task_type", "document")
    iteration = state.get("iteration", 1)
    max_iterations = state.get("max_iterations", 6 if task_type == "coding" else (3 if task_type == "vision" else 4))
    current_step_index = state.get("current_step_index", 0)
    plan = state.get("plan", [])
    observations = state.get("observations", [])

    # Check for any fatal errors in observations
    has_errors = any(obs.get("level") == "error" for obs in observations)
    validation_passed = not has_errors
    artifact_id: Any = state.get("final_artifact_id")
    validation_notes: str = ""

    # Infrastructure/tool execution failure short-circuit (TRD Section 20, ADR-008):
    # a sandbox/Docker failure means pytest never executed, so it must never be
    # re-planned as a test failure or converted into a successful outcome.
    # Terminate explicitly with status="failed" instead of looping.
    latest_struct = observations[-1].get("structured_data", {}) if observations else {}
    if latest_struct.get("execution_status") == "error":
        error_type = latest_struct.get("error_type", "tool_error")
        error_message = latest_struct.get("error_message", "unknown tool execution error")
        validation_notes = f"Infrastructure/tool execution failure ({error_type}): {error_message}"
        broadcaster = get_event_broadcaster()
        logger.error(f"[{task_id}] Validation aborted: {validation_notes}")
        broadcaster.log_and_emit(
            task_id=task_id,
            node="validation",
            message=f"Infrastructure failure ({error_type}): {error_message}. Terminating without self-correction.",
            level="error",
        )
        return {
            "validation_passed": False,
            "status": "failed",
            "validation_notes": validation_notes,
            "final_artifact_id": artifact_id,
        }

    # Task-specific validation logic
    if task_type == "coding":
        if current_step_index < len(plan):
            if has_errors:
                validation_passed = False
                validation_notes = "Tool execution encountered fatal errors."
            else:
                validation_passed = True
                validation_notes = "Plan step executed successfully. Proceeding to next step."
        else:
            latest_test_obs = None
            for obs in reversed(observations):
                struct = obs.get("structured_data", {})
                if struct.get("tool_name") == "run_tests":
                    latest_test_obs = struct
                    break

            if latest_test_obs is not None:
                passed = latest_test_obs.get("passed", False)
                if not passed:
                    validation_passed = False
                    tf = latest_test_obs.get("test_failure", {})
                    failing = tf.get("failing_tests", [])
                    err_summary = tf.get("error_summary", "Tests failed")
                    validation_notes = f"Test failure in {failing}: {err_summary}"
                else:
                    validation_passed = True
                    validation_notes = "All test cases passed in isolated Docker sandbox."
            elif has_errors:
                validation_passed = False
                validation_notes = "Tool execution encountered fatal errors."
            else:
                validation_passed = True
                validation_notes = "Coding workflow completed successfully."

    elif task_type == "document":
        if current_step_index < len(plan):
            if has_errors:
                validation_passed = False
                validation_notes = "Tool execution encountered fatal errors."
            else:
                validation_passed = True
                validation_notes = "Plan step executed successfully. Proceeding to next step."
        else:
            if validation_passed:
                with get_db_context() as session:
                    repo = ArtifactRepository(session)
                    arts = repo.list_by_task_id(task_id)
                    if arts:
                        latest_art = arts[-1]
                        artifact_id = latest_art.artifact_id
                        file_path = Path(latest_art.storage_path)
                        doc_kind = (latest_art.kind or "").lower()
                        
                        if not file_path.exists():
                            validation_passed = False
                            validation_notes = f"Generated {doc_kind.upper()} artifact missing from disk: {file_path}"
                        else:
                            try:
                                if doc_kind == "docx" or file_path.suffix.lower() == ".docx":
                                    doc = docx.Document(str(file_path))
                                    full_doc_text = "\n".join(p.text for p in doc.paragraphs)
                                    doc_text_lower = full_doc_text.lower()

                                    if is_approval_note_goal(state.get("goal") or ""):
                                        required_markers = ["Inspection", "Critical", "Compliance", "Recommendation"]
                                        forbidden_markers: List[str] = []
                                    else:
                                        # Structured analysis report: the sections requested from
                                        # the uploaded document plus source attribution.
                                        required_markers = [
                                            "Main Topic", "Objectives", "Methodology",
                                            "Key Findings", "Conclusions", "Sources",
                                        ]
                                        # Anti-hallucination guard: canned industrial demo content
                                        # must never appear in a report grounded in an unrelated
                                        # uploaded document.
                                        forbidden_markers = [
                                            "Technical Approval Note", "PRV-204",
                                            "Refining Unit 02", "P-102A", "MRPL-INSP",
                                        ]

                                    missing = [m for m in required_markers if m.lower() not in doc_text_lower]
                                    present_forbidden = [m for m in forbidden_markers if m.lower() in doc_text_lower]
                                    if missing:
                                        validation_passed = False
                                        validation_notes = f"DOCX artifact missing required sections: {missing}"
                                    elif present_forbidden:
                                        validation_passed = False
                                        validation_notes = (
                                            f"DOCX artifact contains ungrounded demo content: {present_forbidden}. "
                                            "Report content must be grounded in the uploaded source document."
                                        )
                                    else:
                                        validation_notes = "Plan execution and deliverable artifacts validated successfully."

                                elif doc_kind == "xlsx" or file_path.suffix.lower() == ".xlsx":
                                    wb = openpyxl.load_workbook(str(file_path), read_only=True)
                                    sheets = wb.sheetnames
                                    if len(sheets) < 1:
                                        validation_passed = False
                                        validation_notes = "XLSX artifact contains no worksheets."
                                    else:
                                        validation_notes = f"XLSX artifact validated successfully ({len(sheets)} worksheets: {sheets})."
                                    wb.close()

                                elif doc_kind == "pptx" or file_path.suffix.lower() == ".pptx":
                                    prs = pptx.Presentation(str(file_path))
                                    slide_count = len(prs.slides)
                                    if slide_count < 3:
                                        validation_passed = False
                                        validation_notes = f"PPTX artifact has insufficient slides ({slide_count} < 3)."
                                    else:
                                        validation_notes = f"PPTX presentation validated successfully ({slide_count} slides)."

                                elif doc_kind == "pdf" or file_path.suffix.lower() == ".pdf":
                                    pdf_doc = fitz.open(str(file_path))
                                    page_count = len(pdf_doc)
                                    if page_count < 1:
                                        validation_passed = False
                                        validation_notes = "PDF artifact contains 0 pages."
                                    else:
                                        pdf_text = "\n".join(page.get_text() for page in pdf_doc)
                                        if len(pdf_text.strip()) < 50:
                                            validation_passed = False
                                            validation_notes = "PDF artifact contains insufficient text content."
                                        else:
                                            validation_notes = f"PDF report validated successfully ({page_count} pages, {len(pdf_text)} chars)."
                                    pdf_doc.close()

                                else:
                                    validation_notes = f"Document artifact '{file_path.name}' verified on disk."

                            except Exception as e:
                                validation_passed = False
                                validation_notes = f"Document artifact verification failed for {doc_kind}: {e}"

    elif task_type == "vision":
        # Multimodal Vision Validation (TRD §18.1, Table 48)
        if current_step_index < len(plan):
            if has_errors:
                validation_passed = False
                validation_notes = "Vision step encountered fatal errors."
            else:
                validation_passed = True
                validation_notes = "Plan step executed successfully. Proceeding to next step."
        else:
            latest_vision_res = None
            for obs in reversed(observations):
                struct = obs.get("structured_data", {})
                if struct.get("type") == "vision" and struct.get("vision_result"):
                    latest_vision_res = struct.get("vision_result")
                    break

            if latest_vision_res:
                obs_list = latest_vision_res.get("observation", [])
                interp_list = latest_vision_res.get("interpretation", [])
                uncert_list = latest_vision_res.get("uncertainty", [])
                model_used = latest_vision_res.get("model_used")

                if not obs_list or len(obs_list) == 0:
                    validation_passed = False
                    validation_notes = "VisionResult validation failed: 'observation' array is empty."
                elif not isinstance(interp_list, list) or not isinstance(uncert_list, list):
                    validation_passed = False
                    validation_notes = "VisionResult validation failed: interpretation or uncertainty is not an array."
                elif not model_used:
                    validation_passed = False
                    validation_notes = "VisionResult validation failed: 'model_used' is missing."
                else:
                    validation_passed = True
                    validation_notes = f"Vision findings validated successfully ({len(obs_list)} observations, {len(interp_list)} interpretations, {len(uncert_list)} uncertainties)."
            elif has_errors:
                validation_passed = False
                validation_notes = "Vision workflow encountered execution errors."
            else:
                validation_passed = True
                validation_notes = "Vision workflow completed successfully."

    broadcaster = get_event_broadcaster()
    logger.info(f"[{task_id}] Validation check: passed={validation_passed} (iteration {iteration}/{max_iterations}, step {current_step_index}/{len(plan)}, task_type: {task_type})")

    if validation_passed:
        status = "succeeded" if current_step_index >= len(plan) else "running"
        if not validation_notes:
            validation_notes = "Plan execution validated successfully."
        broadcaster.log_and_emit(
            task_id=task_id,
            node="validation",
            message=f"Validation passed on iteration {iteration}/{max_iterations} (step {current_step_index}/{len(plan)}).",
            level="info",
        )
    else:
        if iteration < max_iterations:
            status = "running"
            if not validation_notes:
                validation_notes = f"Step failed on iteration {iteration}. Self-correction re-plan triggered."
            broadcaster.log_and_emit(
                task_id=task_id,
                node="validation",
                message=f"Validation failed (iteration {iteration}/{max_iterations}). Re-planning...",
                level="warn",
            )
        else:
            status = "failed_bounded"
            validation_notes = f"Max iterations ({max_iterations}) reached without successful validation. Last error: {validation_notes}"
            broadcaster.log_and_emit(
                task_id=task_id,
                node="validation",
                message=f"Iteration limit reached ({max_iterations}/{max_iterations}). Terminating with status=failed_bounded.",
                level="error",
            )

    return {
        "validation_passed": validation_passed,
        "status": status,
        "validation_notes": validation_notes,
        "final_artifact_id": artifact_id,
    }
